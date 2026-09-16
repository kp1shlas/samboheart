"""
Сервис формирования отчётов владельца.
"""
from django.db.models import Count, Sum, Q, Coalesce

from core.models import Child, Attendance, Payment


def build_report(date_from, date_to, group_id=None):
    """
    Отчёт по АКТИВНЫМ детям за период.
    Если group_id указан — только дети с активной записью в эту группу.
    """
    children_qs = Child.objects.filter(is_active=True)

    if group_id:
        children_qs = children_qs.filter(
            enrollments__group_id=group_id,
            enrollments__is_active=True,
        ).distinct()

    rows = []

    for child in children_qs.order_by('full_name'):
        # ── Посещаемость за период (без отменённых занятий) ──
        att_qs = Attendance.objects.filter(
            child=child,
            lesson__date__gte=date_from,
            lesson__date__lte=date_to,
            lesson__is_cancelled=False,
        )
        if group_id:
            att_qs = att_qs.filter(lesson__group_id=group_id)

        att = att_qs.aggregate(
            present=Count('id', filter=Q(status='present')),
            absent=Count('id', filter=Q(status='absent')),
            excused=Count('id', filter=Q(status='excused')),
            excused_reason=Count('id', filter=Q(status='excused_reason')),
            total=Count('id'),
        )

        # ── Оплаты за период (paid_at или created_at, если paid_at пуст) ──
        pay_qs = Payment.objects.annotate(
            pay_date=Coalesce('paid_at', 'created_at')
        ).filter(
            child=child,
            status='paid',
            pay_date__date__gte=date_from,
            pay_date__date__lte=date_to,
        )
        if group_id:
            pay_qs = pay_qs.filter(enrollment__group_id=group_id)

        pay = pay_qs.aggregate(
            amount=Sum('amount'),
            lessons=Sum('lessons_count'),
            count=Count('id'),
        )

        # ── Текущий баланс и группы ──
        if group_id:
            enrollments = list(child.enrollments.filter(group_id=group_id, is_active=True))
        else:
            enrollments = list(child.enrollments.filter(is_active=True))

        balance = sum(e.remaining_lessons for e in enrollments)
        groups_names = ', '.join(e.group.name for e in enrollments) or '—'
        is_free = any(e.is_free for e in enrollments)

        visited = att['present'] or 0
        missed_unexcused = att['absent'] or 0
        missed_excused = (att['excused'] or 0) + (att['excused_reason'] or 0)
        total_lessons = att['total'] or 0
        rate = round(visited * 100 / total_lessons) if total_lessons else None

        rows.append({
            'child': child,
            'groups': groups_names,
            'is_free': is_free,
            'visited': visited,
            'missed_unexcused': missed_unexcused,
            'missed_excused': missed_excused,
            'total_lessons': total_lessons,
            'rate': rate,
            'paid_amount': pay['amount'] or 0,
            'paid_lessons': pay['lessons'] or 0,
            'payments_count': pay['count'] or 0,
            'balance': balance,
            'debt': abs(balance) if balance < 0 else 0,
        })

    totals = {
        'children_count': len(rows),
        'visited': sum(r['visited'] for r in rows),
        'missed_unexcused': sum(r['missed_unexcused'] for r in rows),
        'missed_excused': sum(r['missed_excused'] for r in rows),
        'total_lessons': sum(r['total_lessons'] for r in rows),
        'paid_amount': sum(r['paid_amount'] for r in rows),
        'paid_lessons': sum(r['paid_lessons'] for r in rows),
        'payments_count': sum(r['payments_count'] for r in rows),
        'debt': sum(r['debt'] for r in rows),
    }

    return {
        'rows': rows,
        'totals': totals,
        'date_from': date_from,
        'date_to': date_to,
        'group_id': group_id,
    }


def report_to_excel(report, group_name='Все группы'):
    """Экспорт отчёта в Excel: лист «Отчёт» + лист «Оплаты за период»."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from django.http import HttpResponse

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Отчёт'

    header_fill = PatternFill('solid', start_color='C00000')
    header_font = Font(bold=True, color='FFFFFF')
    total_fill = PatternFill('solid', start_color='F2F2F2')
    thin = Side(style='thin', color='DDDDDD')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws['A1'] = 'Отчёт по занятиям и оплатам'
    ws['A1'].font = Font(bold=True, size=14)
    ws['A2'] = f"Период: {report['date_from']:%d.%m.%Y} — {report['date_to']:%d.%m.%Y}"
    ws['A3'] = f'Группа: {group_name}'
    ws['A4'] = f"Активных детей: {report['totals']['children_count']}"

    headers = [
        'ФИО', 'Группы', 'Посещено', 'Пропущено без ув. причины',
        'Пропущено по ув. причине', 'Всего занятий', 'Посещаемость %',
        'Оплачено ₽', 'Оплачено занятий', 'Платежей', 'Баланс', 'Долг',
    ]
    start_row = 6
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=start_row, column=col, value=h)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(horizontal='center', wrap_text=True)
        c.border = border

    row_i = start_row + 1
    for r in report['rows']:
        values = [
            r['child'].full_name,
            r['groups'] + (' (бесплатно)' if r['is_free'] else ''),
            r['visited'],
            r['missed_unexcused'],
            r['missed_excused'],
            r['total_lessons'],
            f"{r['rate']}%" if r['rate'] is not None else '—',
            float(r['paid_amount']),
            r['paid_lessons'],
            r['payments_count'],
            r['balance'],
            r['debt'],
        ]
        for col, v in enumerate(values, 1):
            c = ws.cell(row=row_i, column=col, value=v)
            c.border = border
        row_i += 1

    t = report['totals']
    total_values = [
        'ИТОГО', f"{t['children_count']} детей",
        t['visited'], t['missed_unexcused'], t['missed_excused'],
        t['total_lessons'],
        f"{round(t['visited'] * 100 / t['total_lessons'])}%" if t['total_lessons'] else '—',
        float(t['paid_amount']), t['paid_lessons'], t['payments_count'],
        '', t['debt'],
    ]
    for col, v in enumerate(total_values, 1):
        c = ws.cell(row=row_i, column=col, value=v)
        c.fill = total_fill
        c.font = Font(bold=True)
        c.border = border

    widths = [32, 26, 10, 16, 16, 12, 14, 12, 14, 10, 9, 8]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = f'A{start_row + 1}'

    # ── Лист 2: детализация оплат ──
    ws2 = wb.create_sheet('Оплаты за период')
    pay_headers = ['Дата', 'Ребёнок', 'Сумма ₽', 'Занятий', 'Способ', 'Комментарий']
    for col, h in enumerate(pay_headers, 1):
        c = ws2.cell(row=1, column=col, value=h)
        c.fill = header_fill
        c.font = header_font

    payments = Payment.objects.annotate(
        pay_date=Coalesce('paid_at', 'created_at')
    ).filter(
        status='paid',
        pay_date__date__gte=report['date_from'],
        pay_date__date__lte=report['date_to'],
        child__is_active=True,
    ).select_related('child', 'enrollment__group').order_by('pay_date')
    if report['group_id']:
        payments = payments.filter(enrollment__group_id=report['group_id'])

    r_i = 2
    for p in payments:
        ws2.cell(row=r_i, column=1, value=p.pay_date.strftime('%d.%m.%Y %H:%M'))
        ws2.cell(row=r_i, column=2, value=p.child.full_name)
        ws2.cell(row=r_i, column=3, value=float(p.amount))
        ws2.cell(row=r_i, column=4, value=p.lessons_count)
        ws2.cell(row=r_i, column=5, value='онлайн' if p.method == 'online' else 'наличные')
        ws2.cell(row=r_i, column=6, value=p.note)
        r_i += 1

    for i, w in enumerate([18, 32, 12, 10, 12, 50], 1):
        ws2.column_dimensions[get_column_letter(i)].width = w
    ws2.freeze_panes = 'A2'

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    filename = f"report_{report['date_from']:%Y%m%d}_{report['date_to']:%Y%m%d}.xlsx"
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response